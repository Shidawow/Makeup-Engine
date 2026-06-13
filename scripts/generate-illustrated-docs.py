#!/usr/bin/env python3
"""Generate illustrated Makeup Engine design and operation DOCX deliverables.

This script is documentation-only. It creates Markdown source files, lightweight
SVG diagram assets, PNG renderings for Word embedding, and two DOCX files on the
Mac desktop. It does not modify application code or require runtime dependencies.
"""

from __future__ import annotations

from dataclasses import dataclass
from html import escape
import os
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
REPORT_MD = ROOT / "docs/reports/makeup-engine-illustrated-detailed-design.md"
MANUAL_MD = ROOT / "docs/manuals/makeup-engine-illustrated-operation-manual.md"
DIAGRAM_DIR = ROOT / "docs/assets/diagrams"
MANUAL_ASSET_DIR = ROOT / "docs/assets/manual"
EXPORT_DIR = ROOT / "tmp/docs-export"
DESKTOP = Path("/Users/star/Desktop")
REPORT_DOCX = DESKTOP / "Makeup-Engine-图文版详细设计书.docx"
MANUAL_DOCX = DESKTOP / "Makeup-Engine-图文版操作说明书.docx"
DELIVERY_NOTE = DESKTOP / "Makeup-Engine-文档交付说明.md"


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "5A6673"
LIGHT = "F4F7FB"
LINE = "C8D3E0"
GREEN = "2E7D5F"
GOLD = "8A6D1D"
RED = "A33A3A"


@dataclass(frozen=True)
class Diagram:
    slug: str
    title: str
    nodes: tuple[str, ...]
    directory: Path
    note: str = ""

    @property
    def svg_path(self) -> Path:
        return self.directory / f"{self.slug}.svg"

    @property
    def png_path(self) -> Path:
        return EXPORT_DIR / f"{self.slug}.png"


def ensure_dirs() -> None:
    for path in [REPORT_MD.parent, MANUAL_MD.parent, DIAGRAM_DIR, MANUAL_ASSET_DIR, EXPORT_DIR, DESKTOP]:
        path.mkdir(parents=True, exist_ok=True)


def font(size: int, bold: bool = False) -> ImageFont.ImageFont:
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Helvetica.ttc",
        "/Library/Fonts/Arial Unicode.ttf",
    ]
    for candidate in candidates:
        try:
            return ImageFont.truetype(candidate, size)
        except Exception:
            continue
    return ImageFont.load_default()


def wrap_text(text: str, width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        trial = f"{current} {word}".strip()
        if len(trial) <= width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines or [text]


def draw_svg(diagram: Diagram) -> None:
    width = 920
    height = 150 + len(diagram.nodes) * 92
    box_x = 110
    box_w = 700
    box_h = 54
    parts = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
        "<defs>",
        '<marker id="arrow" markerWidth="10" markerHeight="10" refX="8" refY="3" orient="auto" markerUnits="strokeWidth">',
        '<path d="M0,0 L0,6 L9,3 z" fill="#5A6673"/>',
        "</marker>",
        "</defs>",
        f'<rect x="0" y="0" width="{width}" height="{height}" rx="18" fill="#F8FAFD"/>',
        f'<text x="46" y="48" font-family="Arial, Helvetica, sans-serif" font-size="24" font-weight="700" fill="#{INK}">{escape(diagram.title)}</text>',
    ]
    y = 90
    colors = [BLUE, DARK_BLUE, GREEN, GOLD, MUTED]
    for index, node in enumerate(diagram.nodes):
        color = colors[index % len(colors)]
        parts.append(f'<rect x="{box_x}" y="{y}" width="{box_w}" height="{box_h}" rx="10" fill="#FFFFFF" stroke="#{color}" stroke-width="2"/>')
        label_lines = wrap_text(node, 54)
        first_y = y + 25 - max(0, len(label_lines) - 1) * 9
        for offset, line in enumerate(label_lines):
            parts.append(
                f'<text x="{box_x + 22}" y="{first_y + offset * 18}" font-family="Arial, Helvetica, sans-serif" '
                f'font-size="15" font-weight="600" fill="#{INK}">{escape(line)}</text>'
            )
        if index < len(diagram.nodes) - 1:
            xmid = box_x + box_w / 2
            parts.append(
                f'<line x1="{xmid}" y1="{y + box_h + 8}" x2="{xmid}" y2="{y + 82}" '
                f'stroke="#{MUTED}" stroke-width="2" marker-end="url(#arrow)"/>'
            )
        y += 92
    if diagram.note:
        parts.append(
            f'<text x="46" y="{height - 28}" font-family="Arial, Helvetica, sans-serif" font-size="13" fill="#{MUTED}">'
            f'{escape(diagram.note)}</text>'
        )
    parts.append("</svg>")
    diagram.svg_path.write_text("\n".join(parts), encoding="utf-8")


def draw_png(diagram: Diagram) -> None:
    width = 1380
    height = 230 + len(diagram.nodes) * 138
    image = Image.new("RGB", (width, height), "#F8FAFD")
    draw = ImageDraw.Draw(image)
    title_font = font(36, True)
    body_font = font(24, True)
    note_font = font(18)
    draw.rounded_rectangle((0, 0, width - 1, height - 1), radius=28, fill="#F8FAFD", outline="#E2E8F0", width=2)
    draw.text((70, 66), diagram.title, fill=f"#{INK}", font=title_font)
    box_x = 165
    box_w = 1050
    box_h = 84
    y = 138
    colors = [BLUE, DARK_BLUE, GREEN, GOLD, MUTED]
    for index, node in enumerate(diagram.nodes):
        color = colors[index % len(colors)]
        draw.rounded_rectangle((box_x, y, box_x + box_w, y + box_h), radius=18, fill="#FFFFFF", outline=f"#{color}", width=4)
        lines = wrap_text(node, 54)
        first_y = y + 31 - max(0, len(lines) - 1) * 14
        for offset, line in enumerate(lines):
            draw.text((box_x + 34, first_y + offset * 28), line, fill=f"#{INK}", font=body_font)
        if index < len(diagram.nodes) - 1:
            xmid = box_x + box_w // 2
            draw.line((xmid, y + box_h + 14, xmid, y + 124), fill=f"#{MUTED}", width=4)
            draw.polygon([(xmid - 12, y + 124), (xmid + 12, y + 124), (xmid, y + 144)], fill=f"#{MUTED}")
        y += 138
    if diagram.note:
        draw.text((70, height - 54), diagram.note, fill=f"#{MUTED}", font=note_font)
    image.save(diagram.png_path)


def diagrams() -> dict[str, Diagram]:
    specs = [
        Diagram(
            "project-position",
            "Project Position",
            (
                "Makeup Engine: template production system",
                "User App Shell: local contract-driven PWA prototype",
                "Future Production App: separate app after explicit gate",
            ),
            DIAGRAM_DIR,
        ),
        Diagram(
            "system-architecture",
            "System Architecture",
            (
                "Vision subsystem and source artifact binding",
                "Template Studio admin workbench",
                "Template Library and Publish Package",
                "UserAppTemplatePackage contract",
                "User App Shell and internal trial admin panels",
                "Docs, project-state, recovery, provider handoff",
            ),
            DIAGRAM_DIR,
        ),
        Diagram(
            "end-to-end-data-flow",
            "End-to-End Data Flow",
            (
                "Source image or manual template input",
                "Vision analysis, regions, segmentation foundation",
                "Weighted sampling and makeup semantics",
                "Template builder, correction, QA evidence",
                "Template Library and UserAppTemplatePackage",
                "User App Shell, trial pack, evidence review, follow-up",
            ),
            DIAGRAM_DIR,
        ),
        Diagram(
            "objective-roadmap",
            "Objective Roadmap",
            (
                "Explainable template production",
                "Auditable correction and evidence",
                "Validated UserAppTemplatePackage",
                "Local PWA / Mobile Web MVP shell",
                "Anonymous internal trial learning",
                "Future validation plan after explicit gate",
            ),
            DIAGRAM_DIR,
        ),
        Diagram(
            "vision-pipeline",
            "Vision Pipeline",
            (
                "SourceImagePackage manifest",
                "Operator artifact binding",
                "Face detection and landmarks",
                "Cosmetic regions and mask foundation",
                "Weighted pixel sampling and baseline analysis",
                "Correction persistence and evidence",
            ),
            DIAGRAM_DIR,
        ),
        Diagram(
            "mask-weighted-sampling",
            "Mask Correction and Weighted Sampling",
            (
                "Cosmetic region candidate mask",
                "Operator review and editable correction",
                "Weighted pixel sampling",
                "Skin baseline and edge analysis",
                "Template evidence and review queue",
            ),
            DIAGRAM_DIR,
        ),
        Diagram(
            "template-studio-workflow",
            "Template Studio Workflow",
            (
                "Source Image Intake",
                "Vision Analysis and Overlay Debug",
                "Editable Mask Workbench",
                "Evidence, Dataset, Review Queue",
                "Template Library and Publish Package",
                "Package Preview and Prototype Consumer",
            ),
            DIAGRAM_DIR,
        ),
        Diagram(
            "contract-boundary",
            "Contract Boundary",
            (
                "SourceImagePackage stays in Makeup Engine",
                "Correction, evidence, review, package validation",
                "TemplatePublishPackage creates local export metadata",
                "UserAppTemplatePackage becomes consumption contract",
                "User App Shell consumes contract only",
            ),
            DIAGRAM_DIR,
        ),
        Diagram(
            "user-vs-admin-path",
            "User Path vs Admin Path",
            (
                "Ordinary user: discover, detail, tools, step guide, privacy",
                "Local-only shell state: preferences and session recovery",
                "Administrator: readiness, QA, trial, evidence, follow-up panels",
                "Admin surfaces do not become production user features",
            ),
            DIAGRAM_DIR,
        ),
        Diagram(
            "internal-trial-loop",
            "Internal Trial Loop: 8C to 9J",
            (
                "8C Trial Pack and feedback preview",
                "8D Content QA and 8E readiness gate",
                "9A Operations and 9B result review",
                "9C Iteration, 9D learning, 9E evidence",
                "9F collection prep, 9G dry run, 9H launch",
                "9I evidence review and 9J follow-up iteration",
            ),
            DIAGRAM_DIR,
        ),
        Diagram(
            "evidence-lifecycle",
            "Anonymous Evidence Lifecycle",
            (
                "Trial task and participant notice",
                "Anonymous observation notes",
                "Evidence pack and sufficiency gate",
                "Dry run, launch, post-launch handoff",
                "Evidence review, gap review, decision input",
                "Follow-up iteration and gap action plan",
            ),
            DIAGRAM_DIR,
        ),
        Diagram(
            "go-no-go-decision",
            "Go / No-Go Decision Logic",
            (
                "Privacy or forbidden-data issue",
                "Pause and fix scope before continuing",
                "Content-heavy issue",
                "Revise template content before more trials",
                "Shell-heavy issue",
                "Revise shell or continue with warnings",
            ),
            DIAGRAM_DIR,
        ),
        Diagram(
            "privacy-boundary",
            "Privacy Boundary",
            (
                "Allowed: anonymous notes and aggregated issue counts",
                "Forbidden: photos, names, contact, health, sensitive identity",
                "Forbidden: image bytes, base64, object URLs, local paths",
                "No backend, no analytics, no OpenAI/external API, no training",
                "No real user trial records in project-state",
            ),
            DIAGRAM_DIR,
        ),
        Diagram(
            "forbidden-data-flow",
            "Forbidden Data Flow",
            (
                "Real photo, name, contact, health, identity, biometrics",
                "Blocked before shell state, project-state, backend, analytics",
                "No upload, no training dataset, no AI analysis",
                "Stop trial recording and fix privacy scope",
            ),
            DIAGRAM_DIR,
        ),
        Diagram(
            "qa-gate",
            "QA Gate Flow",
            (
                "Typecheck",
                "Scoped unit and component tests",
                "Full test suite",
                "Build",
                "Project status and context scripts",
                "Documentation and project-state recovery tests",
            ),
            DIAGRAM_DIR,
        ),
        Diagram(
            "recovery-flow",
            "Provider Handoff and Recovery",
            (
                "START_HERE",
                "docs/status and NEXT_ACTION",
                "MASTER_CODEX_CONTEXT and PROVIDER_SWITCH_PROMPT",
                "project-state snapshot and handoffs",
                "Active task, guardrails, command log",
                "Git branch, commit, push",
            ),
            DIAGRAM_DIR,
        ),
        Diagram(
            "startup-flow",
            "Local Startup Flow",
            (
                "cd /Users/star/Makeup-Engine",
                "npm install",
                "npm run dev",
                "Open local Vite URL",
                "Run status, context, typecheck, build, test",
            ),
            MANUAL_ASSET_DIR,
        ),
        Diagram(
            "branch-workflow",
            "Git Branch Workflow",
            (
                "Create phase or docs branch",
                "Make scoped docs/code/test changes",
                "Run validation commands",
                "Stage explicit files only",
                "Commit and push to GitHub",
            ),
            MANUAL_ASSET_DIR,
        ),
        Diagram(
            "status-recovery-flow",
            "Status Recovery Flow",
            (
                "Read START_HERE",
                "Read CURRENT_PROJECT_STATUS and NEXT_ACTION",
                "Read project-state snapshot and handoffs",
                "Run project:status and project:context",
                "Resume from repository source of truth",
            ),
            MANUAL_ASSET_DIR,
        ),
        Diagram(
            "template-studio-wireframe",
            "Template Studio Page Structure",
            (
                "Left: source image and task selection",
                "Center: overlay, mask workbench, replay view",
                "Right: evidence, dataset, review queue",
                "Bottom: library, publish package, preview, prototype consumer",
            ),
            MANUAL_ASSET_DIR,
        ),
        Diagram(
            "ordinary-user-flow",
            "Ordinary User Flow",
            (
                "Discover makeup templates",
                "Open template detail",
                "Prepare tools and products",
                "Start guided steps",
                "Review local progress, preferences, privacy notice",
            ),
            MANUAL_ASSET_DIR,
        ),
        Diagram(
            "admin-overview",
            "Administrator Panels",
            (
                "PWA, MVP polish, readiness, mobile QA",
                "Trial pack, feedback, content QA, release gate",
                "Trial ops, result review, iteration, learning gate",
                "Evidence pack, collection prep, dry run, launch",
                "Evidence review, gap review, follow-up readiness",
            ),
            MANUAL_ASSET_DIR,
        ),
        Diagram(
            "trial-prep-flow",
            "Anonymous Trial Preparation",
            (
                "Confirm readiness gate",
                "Check evidence collection protocol",
                "Run dry run pack",
                "Prepare launch pack, participant notice, admin script",
                "Confirm allowed evidence, forbidden data, stop conditions",
            ),
            MANUAL_ASSET_DIR,
        ),
        Diagram(
            "trial-review-tree",
            "Anonymous Trial Review Decision Tree",
            (
                "Review anonymous evidence",
                "Identify evidence gaps",
                "Build decision input",
                "Create follow-up iteration and action plan",
                "Continue, repeat dry run, revise, pause, or prepare 10A preconditions",
            ),
            MANUAL_ASSET_DIR,
        ),
    ]
    return {diagram.slug: diagram for diagram in specs}


def generate_diagrams(diagram_map: dict[str, Diagram]) -> None:
    for diagram in diagram_map.values():
        draw_svg(diagram)
        draw_png(diagram)


def image_ref(diagram: Diagram, base_dir: Path) -> str:
    return f"![{diagram.title}]({os.path.relpath(diagram.svg_path, base_dir).replace(os.sep, '/')})"


def md_section(title: str, body: str) -> str:
    return f"## {title}\n\n{body.strip()}\n"


def design_markdown(diagram_map: dict[str, Diagram]) -> str:
    d = diagram_map
    base_dir = REPORT_MD.parent
    screenshot_note = (
        "> 本次文档未采集真实界面截图，使用架构图、流程图和页面结构示意图替代。"
        "后续可在运行环境稳定后补充截图。"
    )
    sections = [
        md_section(
            "1. 项目概述",
            f"""
Makeup Engine 是妆容模板生产系统，用于把可审核的视觉分析、人工修正、模板证据、模板库和发布包组织成未来用户侧 App 可以消费的模板契约。User App Shell 是本地 contract-driven PWA / Mobile Web prototype，用于验证 UserAppTemplatePackage 的消费路径、移动端体验和匿名内部试用准备。

当前路线是 React Web / PWA MVP first。当前不是 production app，也不做 iOS 原生、React Native、Flutter、后端、数据库、账号、相机、AR、OpenAI API、训练或 App Store 发布。

{image_ref(d['project-position'], base_dir)}
""",
        ),
        md_section(
            "2. 总体目标",
            f"""
总体目标是建立可解释、可审核、可发布的妆容模板生产体系，并用 UserAppTemplatePackage 作为未来用户侧 App 的消费契约。当前能力覆盖 PWA / Mobile Web MVP Shell、匿名内部试用准备、证据复盘和 Phase 9J 后续迭代规划。

关键目标：

- 模板生产链路可追溯：来源、分析、修正、证据、QA、发布包均可审计。
- 用户侧消费契约稳定：UserAppTemplatePackage 经过本地 JSON round-trip 和 compatibility validation。
- 试用决策保守：内部试用只使用匿名、本地、示例级或安全内部证据，不直接推导生产发布。

{image_ref(d['objective-roadmap'], base_dir)}
""",
        ),
        md_section(
            "3. 系统总体架构",
            f"""
系统由 Vision subsystem、Template Studio、Template Library、Template Production Batch、Template Publish Package、UserAppTemplatePackage、User App Shell prototype、Trial Pack、Content QA、Release Readiness Gate、Internal Trial Operations、Evidence Collection、Dry Run、Launch Pack、Evidence Review、Follow-up Iteration、Project state / recovery docs 组成。

{image_ref(d['system-architecture'], base_dir)}

模块关系的核心原则是：Makeup Engine 负责模板生产和本地验证；User App Shell 负责读取契约并模拟未来 PWA/Mobile Web 用户体验；未来 production app 必须等待新的明确 phase gate。
""",
        ),
        md_section(
            "4. 核心数据流",
            f"""
Source image / manual template input 进入 vision analysis 后，产生 cosmetic regions、segmentation foundation、weighted sampling、makeup parameters 和 semantics。随后 template builder 生成模板草案，经过 template QA、correction、evidence、Template Library，再导出 UserAppTemplatePackage，最后由 User App Shell 消费，并进入 trial pack、feedback、content QA、release readiness、trial ops、evidence、dry run、launch、review 和 follow-up。

关键边界：

- SourceImagePackage 不能直接进入 User App。
- SourceImagePackage 不是 training dataset。
- UserAppTemplatePackage 是消费契约，不是在线发布物。
- User App Shell 不能直接读取 SourceImagePackage。

{image_ref(d['end-to-end-data-flow'], base_dir)}

{image_ref(d['contract-boundary'], base_dir)}
""",
        ),
        md_section(
            "5. Vision-first 管线设计",
            f"""
Vision-first 管线覆盖 Face detection / landmarks、MediaPipe FaceMesh provider、cosmetic regions、segmentation foundation、weighted pixel sampling、skin baseline / edge analysis、editable mask、correction persistence、dataset export / review queue、PNG image codec / artifact manifest。

当前限制包括：未接真实深度学习 segmentation model；JPEG decode intentionally unsupported；`public/mediapipe` 需要单独准备且不提交；当前主要是本地分析和可审核证据链。

{image_ref(d['vision-pipeline'], base_dir)}

{image_ref(d['mask-weighted-sampling'], base_dir)}
""",
        ),
        md_section(
            "6. Template Studio 设计",
            f"""
Template Studio 是本地管理员工作台，覆盖 Source Image Intake、Vision Analysis、Overlay Debug、Editable Mask Workbench、Evidence Panel、Dataset Panel、Review Queue、Replay Viewer、Template Library、Publish Package、Package Preview 和 User App Prototype Consumer。

{image_ref(d['template-studio-workflow'], base_dir)}
""",
        ),
        md_section(
            "7. Template / App Contract 设计",
            f"""
核心 contract 包括 MakeupTemplate、TemplateEvidence、HumanCorrectionDataset、TemplateLibraryEntry、TemplatePublishPackage、UserAppTemplatePackage。契约层要求 compatibility validation、durable export 不含 object URL / local path / image bytes / React state，并保持 JSON round-trip stability。

UserAppTemplatePackage 是 Makeup Engine 与未来用户侧 App 的消费边界。它可以被 User App Shell 读取，但不能反向修改模板生产状态。

{image_ref(d['contract-boundary'], base_dir)}
""",
        ),
        md_section(
            "8. User App Shell 设计",
            f"""
User App Shell 是 local PWA / Mobile Web MVP Shell，覆盖 template list / detail、step-by-step guidance、tools / product suggestions、local onboarding、local preferences、local session persistence、discovery / recommendation placeholder、photo / personalization placeholder、privacy notice、PWA readiness、mobile web polish、browser/mobile QA、trial pack、content QA、release readiness gate、internal trial ops、evidence collection、dry run、launch pack、evidence review 和 follow-up iteration。

{image_ref(d['user-vs-admin-path'], base_dir)}
""",
        ),
        md_section(
            "9. 匿名内部试用体系设计",
            f"""
内部试用体系从 Phase 8C 到 Phase 9J 逐步形成闭环：8C User App MVP Trial Pack、8D Template Content QA、8E MVP Release Readiness Gate、9A Internal Trial Operations、9B Trial Result Review、9C Trial Iteration Plan、9D Learning Summary & Decision Gate、9E Internal Trial Evidence Pack、9F Evidence Collection Preparation、9G Anonymous Dry Run Pack、9H Anonymous Trial Launch Pack、9I Anonymous Trial Evidence Review、9J Follow-up Iteration。

{image_ref(d['internal-trial-loop'], base_dir)}

{image_ref(d['evidence-lifecycle'], base_dir)}

{image_ref(d['go-no-go-decision'], base_dir)}
""",
        ),
        md_section(
            "10. 隐私与安全边界",
            f"""
隐私边界是当前体系的硬约束：不采集真实照片；不上传照片；不保存 image bytes、base64、object URL、本地路径；不保存 faceEmbedding / biometricId；不收集健康信息、敏感身份信息、精确身份信息；不写真实用户记录到 project-state；不写 trial feedback 到 training dataset；不调用 OpenAI API / external API；不做 backend / database / analytics。

{image_ref(d['privacy-boundary'], base_dir)}

{image_ref(d['forbidden-data-flow'], base_dir)}
""",
        ),
        md_section(
            "11. 当前不做事项",
            """
当前明确不做 production app、iOS native、React Native、Flutter、App Store / TestFlight、backend、database、accounts、cloud sync、camera、AR、OpenAI API、external CV API、training model、ecommerce / payment / community、public launch。
""",
        ),
        md_section(
            "12. 当前已知限制",
            """
已知限制包括：no production user app；no real camera / AR；no backend / sync；`public/mediapipe` intentionally not committed；JPEG decode unsupported；no real deep segmentation model；User App Shell is prototype；readiness gate is not production approval；browser/mobile QA is not real device lab QA；internal trial ops is not public launch；current evidence is anonymous/mock/example/framework level unless manually run in a safe internal setting。
""",
        ),
        md_section(
            "13. 测试与质量保障",
            f"""
质量保障覆盖 typecheck、unit tests、scoped tests、full tests、build、project:status、project:context、context-pack、browser QA、documentation recovery tests、project-state recovery tests。Phase 9J 最近记录为 471 files / 729 tests，Phase 9J scoped tests 为 10 files / 24 tests。

{image_ref(d['qa-gate'], base_dir)}
""",
        ),
        md_section(
            "14. 项目状态恢复机制",
            f"""
恢复机制依赖 START_HERE.md、docs/status、docs/prompts、project-state/*.json、provider handoff、active task、compact handoff 和 GitHub branch / commit workflow。仓库文件是 source of truth，聊天记忆不是。

{image_ref(d['recovery-flow'], base_dir)}
""",
        ),
        md_section(
            "15. 后续路线",
            """
DOC-ILLUSTRATED 本次完成后，可进入 Phase 10A MVP Validation Plan 或按项目节奏先执行 Phase 9K Anonymous Internal Trial Evidence Round 2 Pack。后续仍不默认进入 production app。真实产品验证前还需要明确用户、指标、试用人数、成功标准、失败标准和合规边界。
""",
        ),
    ]
    return "\n\n".join(
        [
            "# Makeup Engine 图文版详细设计书",
            "",
            "版本：Phase 9J 后 | 输出：DOC-ILLUSTRATED | 仓库：https://github.com/Shidawow/Makeup-Engine",
            "",
            screenshot_note,
            "",
            "## 目录",
            "",
            "\n".join(f"- {i}. {title}" for i, title in enumerate([
                "项目概述", "总体目标", "系统总体架构", "核心数据流", "Vision-first 管线设计",
                "Template Studio 设计", "Template / App Contract 设计", "User App Shell 设计",
                "匿名内部试用体系设计", "隐私与安全边界", "当前不做事项", "当前已知限制",
                "测试与质量保障", "项目状态恢复机制", "后续路线"
            ], 1)),
            "",
            *sections,
        ]
    ) + "\n"


def manual_markdown(diagram_map: dict[str, Diagram]) -> str:
    d = diagram_map
    base_dir = MANUAL_MD.parent
    screenshot_note = (
        "> 本次文档未采集真实界面截图，使用架构图、流程图和页面结构示意图替代。"
        "后续可在运行环境稳定后补充截图。"
    )
    admin_items = [
        "PWA 检查", "MVP 打磨", "App 就绪度", "移动端 QA", "交互检查", "MVP 试用包",
        "反馈表预览", "试用就绪度", "模板内容 QA", "试用模板选择", "试用内容就绪度",
        "MVP 发布就绪度", "试用 Go/No-Go", "内部试用运营", "观察记录模板", "试用结果复盘",
        "试用结果复盘框架", "问题分类汇总", "下一步决策框架", "试用迭代计划", "迭代 backlog",
        "优先级建议", "试用学习总结", "产品决策门", "下一阶段建议", "内部试用证据包",
        "试用证据摘要", "证据充分性判断", "证据收集协议", "证据收集 checklist",
        "证据收集质量门", "匿名内部试用 dry run", "dry run checklist", "dry run 复盘",
        "匿名内部试用启动包", "启动就绪度", "试用后 handoff", "匿名试用证据复盘",
        "证据缺口复盘", "下一步决策输入", "匿名试用后续迭代", "证据缺口行动计划", "后续试用就绪度",
    ]
    sections = [
        md_section(
            "1. 这个系统是做什么的",
            f"""
Makeup Engine 是模板生产和试用准备系统。它可以在浏览器本地运行，用于准备 PWA / Mobile Web MVP 和匿名内部试用。它不是正式 App，不是后端服务，也不是公开用户增长系统。

{image_ref(d['project-position'], base_dir)}
""",
        ),
        md_section(
            "2. 如何启动项目",
            f"""
在 MacBook 上固定使用项目目录：

```bash
cd /Users/star/Makeup-Engine
npm install
npm run dev
```

浏览器打开终端显示的本地地址，例如 `http://localhost:5173/`。

常用验证命令：

```bash
npm run project:status
npm run project:context
npm run typecheck
npm run build
npm run test
npm run user-app:browser-qa -- --json
```

{image_ref(d['startup-flow'], base_dir)}
""",
        ),
        md_section(
            "3. 推荐日常工作目录",
            """
日常固定使用 `/Users/star/Makeup-Engine`。不要主要依赖 Codex worktree 临时目录，因为长期状态、桌面交付、GitHub push 和项目恢复都应以这个主目录为准。
""",
        ),
        md_section(
            "4. 分支工作流",
            f"""
每个 Phase 或独立文档任务新建分支，完成后 commit + push GitHub，不直接在 main 上堆功能。不要 `git add -A`，不要提交 `dist`、`node_modules`、`tmp`、`public/mediapipe`、`.DS_Store`。当前 GitHub 仓库是 https://github.com/Shidawow/Makeup-Engine。

{image_ref(d['branch-workflow'], base_dir)}
""",
        ),
        md_section(
            "5. 如何查看当前项目状态",
            f"""
优先查看 `START_HERE.md`、`docs/status/CURRENT_PROJECT_STATUS.md`、`docs/status/NEXT_ACTION.md`、`project-state/project-state.snapshot.json`、`project-state/latest-handoff.json`。

命令：

```bash
npm run project:status
npm run project:context
node scripts/project-status.mjs --json
node scripts/context-pack.mjs --json
```

{image_ref(d['status-recovery-flow'], base_dir)}
""",
        ),
        md_section(
            "6. Template Studio 使用说明",
            f"""
Template Studio 管理员路径包括 Source Image Intake、Vision Analysis、Overlay Debug、Editable Mask Workbench、Evidence Panel、Dataset Panel、Review Queue、Replay Viewer、Template Library、Publish Package、Package Preview、User App Prototype Consumer。入口名称以实际 UI 为准。

{image_ref(d['template-studio-wireframe'], base_dir)}
""",
        ),
        md_section(
            "7. User App Shell 普通用户路径",
            f"""
普通用户路径聚焦：发现妆容、查看妆容详情、准备工具、开始跟练、逐步查看步骤、查看本地进度、查看本地偏好、查看隐私说明。普通用户路径不展示 trial ops、evidence review、follow-up readiness 等管理员术语。

{image_ref(d['ordinary-user-flow'], base_dir)}
""",
        ),
        md_section(
            "8. User App Shell 管理员路径",
            f"""
管理员路径用于 QA、试用准备、证据复盘和后续迭代，不是普通用户功能。当前管理员入口包括：

{chr(10).join(f'- {item}' for item in admin_items)}

{image_ref(d['admin-overview'], base_dir)}

{image_ref(d['user-vs-admin-path'], base_dir)}
""",
        ),
        md_section(
            "9. 如何准备匿名内部试用",
            f"""
准备匿名内部试用时，先确认 readiness gate，再使用 9F evidence collection protocol、9G dry run pack、9H launch pack，并确认 participant notice、admin script、allowed evidence、forbidden data、stop conditions。

{image_ref(d['trial-prep-flow'], base_dir)}
""",
        ),
        md_section(
            "10. 如何复盘匿名内部试用",
            f"""
复盘时使用 9I evidence review、evidence gap review、decision input，再使用 9J follow-up iteration 判断是否继续匿名内部试用、重复 dry run、修 protocol、修 launch pack、暂停或准备 MVP validation preconditions。

{image_ref(d['trial-review-tree'], base_dir)}

{image_ref(d['go-no-go-decision'], base_dir)}
""",
        ),
        md_section(
            "11. 可以做什么 / 不可以做什么",
            """
可以做：本地运行、查看模板、跟练流程、查看推荐、查看 PWA readiness、做内容 QA、准备匿名内部试用、做匿名证据复盘、做后续迭代计划。

不可以做：正式发布、公开招募、收集真实照片、收集真实姓名 / 手机 / 邮箱、收集健康信息 / 敏感身份 / 生物识别、上线 App Store、接后端、调 OpenAI API、训练模型、收集敏感资料。
""",
        ),
        md_section(
            "12. 常见问题",
            """
**为什么不是 iOS App？** 当前路线是 React Web / PWA MVP first，原生 iOS 等待后续明确 gate。

**为什么先做 PWA？** PWA 能更快验证模板价值、移动端路径和内部试用流程。

**能不能直接给用户用？** 只能用于内部小范围匿名试用准备，不是正式发布。

**能不能收集照片？** 不能。当前 photo 是 placeholder，不调用 camera API。

**能不能接 OpenAI？** 当前不能接 OpenAI API 或外部 AI/CV API。

**public/mediapipe 为什么没有？** 这是 intentionally not committed 的本地资源，需要单独准备。

**build 有 chunk warning 是否严重？** 当前是 Vite 大 chunk 提醒，不等于失败。

**full test 怎么跑？** 使用 `npm run test`。

**Codex worktree 和 /Users/star/Makeup-Engine 有什么区别？** 日常交付以 `/Users/star/Makeup-Engine` 为准。

**如何确认代码已经 push 到 GitHub？** 查看 `git status`、`git branch --show-current`、`git remote -v`，并确认 push 输出成功。

**为什么内部试用不等于正式发布？** 内部试用只验证理解、跟练意愿和模板价值，不开放公开增长。

**什么情况下必须暂停试用？** 发现照片、上传、真实身份、敏感信息、训练或后端记录请求时必须暂停。

**为什么现在不继续堆 Phase 9K？** 当前任务是补正式图文文档和操作说明，先沉淀系统全貌。

**为什么需要先补图文详细设计书和操作说明书？** 方便老板、产品负责人和内部试用管理员理解边界、流程和下一步决策。
""",
        ),
        md_section(
            "13. 后续工作",
            """
DOC-ILLUSTRATED 本次完成后，后续可以进入 Phase 10A MVP Validation Plan 或按证据策略进入 Phase 9K Anonymous Internal Trial Evidence Round 2 Pack。是否继续 production app discovery，应由匿名内部试用证据和明确产品验证指标决定。
""",
        ),
    ]
    return "\n\n".join(
        [
            "# Makeup Engine 图文版操作说明书",
            "",
            "读者：项目老板 / 产品负责人 / 内部试用管理员",
            "",
            screenshot_note,
            "",
            "## 目录",
            "",
            "\n".join(f"- {i}. {title}" for i, title in enumerate([
                "这个系统是做什么的", "如何启动项目", "推荐日常工作目录", "分支工作流",
                "如何查看当前项目状态", "Template Studio 使用说明", "User App Shell 普通用户路径",
                "User App Shell 管理员路径", "如何准备匿名内部试用", "如何复盘匿名内部试用",
                "可以做什么 / 不可以做什么", "常见问题", "后续工作"
            ], 1)),
            "",
            *sections,
        ]
    ) + "\n"


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(9.5)
    set_run_font(run)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F4F7" if bold else "FFFFFF")
    tc_pr.append(shd)


def set_run_font(run, name: str = "Calibri", east_asia: str = "PingFang SC") -> None:
    run.font.name = name
    if run._element.rPr is not None:
        run._element.rPr.rFonts.set(qn("w:ascii"), name)
        run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
        run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)


def style_doc(doc: Document, title: str) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
    header_p = section.header.paragraphs[0]
    header_p.text = title
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header_p.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string(MUTED)
        set_run_font(run)
    footer_p = section.footer.paragraphs[0]
    footer_p.text = "Makeup Engine | DOC-ILLUSTRATED"
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer_p.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string(MUTED)
        set_run_font(run)


def add_cover(doc: Document, title: str, subtitle: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("DOC-ILLUSTRATED")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(GOLD)
    set_run_font(run)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor.from_string(INK)
    set_run_font(run)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    run = p.add_run(subtitle)
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    set_run_font(run)

    table = doc.add_table(rows=4, cols=2)
    table.autofit = False
    for row in table.rows:
        row.cells[0].width = Inches(1.65)
        row.cells[1].width = Inches(4.85)
    for label, value, row in [
        ("项目", "Makeup Engine", 0),
        ("阶段", "Phase 9J completed; DOC-ILLUSTRATED task", 1),
        ("仓库", "https://github.com/Shidawow/Makeup-Engine", 2),
        ("边界", "Not production app; no backend/camera/AR/OpenAI/training", 3),
    ]:
        set_cell_text(table.rows[row].cells[0], label, True)
        set_cell_text(table.rows[row].cells[1], value)


def add_callout(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.5)
    set_cell_text(cell, text)
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT)
    tc_pr.append(shd)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_run_font(run)


def add_numbered(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run(text)
    set_run_font(run)


@dataclass(frozen=True)
class DocSection:
    title: str
    paragraphs: tuple[str, ...] = ()
    bullets: tuple[str, ...] = ()
    diagram: str | None = None
    extra_diagram: str | None = None


def add_section(doc: Document, section: DocSection, diagram_map: dict[str, Diagram]) -> None:
    doc.add_heading(section.title, level=1)
    for paragraph in section.paragraphs:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.25
        run = p.add_run(paragraph)
        set_run_font(run)
    for bullet in section.bullets:
        add_bullet(doc, bullet)
    for slug in [section.diagram, section.extra_diagram]:
        if slug:
            diagram = diagram_map[slug]
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            run.add_picture(str(diagram.png_path), width=Inches(5.9))
            caption = doc.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.paragraph_format.space_after = Pt(8)
            cap = caption.add_run(f"图：{diagram.title}")
            cap.italic = True
            cap.font.size = Pt(9)
            cap.font.color.rgb = RGBColor.from_string(MUTED)
            set_run_font(cap)


def design_sections() -> list[DocSection]:
    return [
        DocSection("1. 项目概述", ("Makeup Engine 是妆容模板生产系统；User App Shell 是本地 contract-driven PWA / Mobile Web prototype；当前路线是 React Web / PWA MVP first，但当前不是 production app。",), ("不做 iOS 原生、React Native、Flutter、后端、数据库、账号、相机、AR、OpenAI API、训练和 App Store。",), "project-position"),
        DocSection("2. 总体目标", ("目标是建立可解释、可审核、可发布的妆容模板，并用 UserAppTemplatePackage 供未来用户侧 App 消费。当前支持 PWA / Mobile Web MVP Shell 和匿名内部试用准备。",), ("模板生产链路可追溯。", "用户侧契约稳定。", "内部试用决策保持匿名、本地和保守。"), "objective-roadmap"),
        DocSection("3. 系统总体架构", ("系统覆盖 Vision subsystem、Template Studio、Template Library、Template Production Batch、Template Publish Package、UserAppTemplatePackage、User App Shell prototype、Trial Pack、Content QA、Release Readiness Gate、Internal Trial Operations、Evidence Collection、Dry Run、Launch Pack、Evidence Review、Follow-up Iteration、Project state / recovery docs。",), (), "system-architecture"),
        DocSection("4. 核心数据流", ("Source image / manual template input 经过 vision analysis、cosmetic regions、segmentation、weighted sampling、makeup semantics、template builder、QA/correction/evidence、Template Library、UserAppTemplatePackage，最后进入 User App Shell 和 trial/evidence/follow-up 管理闭环。",), ("SourceImagePackage 不能直接进入 User App。", "SourceImagePackage 不是 training dataset。", "UserAppTemplatePackage 是消费契约。", "User App Shell 不能直接读取 SourceImagePackage。"), "end-to-end-data-flow", "contract-boundary"),
        DocSection("5. Vision-first 管线设计", ("Vision-first 管线覆盖 Face detection / landmarks、MediaPipe FaceMesh provider、cosmetic regions、segmentation foundation、weighted pixel sampling、skin baseline / edge analysis、editable mask、correction persistence、dataset export / review queue、PNG image codec / artifact manifest。",), ("当前没有真实深度学习 segmentation model。", "JPEG decode unsupported。", "public/mediapipe 需要单独准备。"), "vision-pipeline", "mask-weighted-sampling"),
        DocSection("6. Template Studio 设计", ("Template Studio 是本地管理员工作台，覆盖 Source Image Intake、Vision Analysis、Overlay Debug、Editable Mask Workbench、Evidence Panel、Dataset Panel、Review Queue、Replay Viewer、Template Library、Publish Package、Package Preview 和 User App Prototype Consumer。",), (), "template-studio-workflow"),
        DocSection("7. Template / App Contract 设计", ("核心对象包括 MakeupTemplate、TemplateEvidence、HumanCorrectionDataset、TemplateLibraryEntry、TemplatePublishPackage、UserAppTemplatePackage。durable export 禁止 object URL、local path、image bytes、React state，并要求 JSON round-trip stability。",), (), "contract-boundary"),
        DocSection("8. User App Shell 设计", ("User App Shell 覆盖 template list/detail、step-by-step guidance、tools/product suggestions、local onboarding、local preferences、session persistence、discovery/recommendation placeholder、photo/personalization placeholder、privacy notice、PWA readiness、mobile polish、browser/mobile QA 和 8C-9J 管理员面板。",), (), "user-vs-admin-path"),
        DocSection("9. 匿名内部试用体系设计", ("匿名内部试用体系覆盖 8C Trial Pack、8D Content QA、8E Release Readiness、9A Ops、9B Result Review、9C Iteration Plan、9D Learning Gate、9E Evidence Pack、9F Collection Prep、9G Dry Run、9H Launch、9I Evidence Review、9J Follow-up Iteration。",), (), "internal-trial-loop", "evidence-lifecycle"),
        DocSection("10. 隐私与安全边界", ("系统不采集真实照片、不上传照片、不保存 image bytes/base64/object URL/local path、不保存 faceEmbedding/biometricId、不收集健康信息/敏感身份/精确身份、不写真实用户记录到 project-state、不写 trial feedback 到 training dataset、不调用 OpenAI API/external API、不做 backend/database/analytics。",), (), "privacy-boundary", "forbidden-data-flow"),
        DocSection("11. 当前不做事项", ("当前不做 production app、iOS native、React Native、Flutter、App Store/TestFlight、backend、database、accounts、cloud sync、camera、AR、OpenAI API、external CV API、training model、ecommerce/payment/community、public launch。",), ()),
        DocSection("12. 当前已知限制", ("已知限制包括 no production user app、no real camera/AR、no backend/sync、public/mediapipe intentionally not committed、JPEG decode unsupported、no real deep segmentation model、User App Shell is prototype、readiness gate is not production approval、browser/mobile QA is not real device lab QA、internal trial ops is not public launch、current evidence is anonymous/mock/example/framework level。",), ()),
        DocSection("13. 测试与质量保障", ("质量保障覆盖 typecheck、unit tests、scoped tests、full tests、build、project:status、project:context、context-pack、browser QA、documentation recovery tests、project-state recovery tests。",), (), "qa-gate"),
        DocSection("14. 项目状态恢复机制", ("恢复机制由 START_HERE.md、docs/status、docs/prompts、project-state/*.json、provider handoff、active task、compact handoff 和 GitHub branch/commit workflow 共同构成。",), (), "recovery-flow"),
        DocSection("15. 后续路线", ("DOC-ILLUSTRATED 本次完成。后续可进入 Phase 10A MVP Validation Plan 或按证据策略执行 Phase 9K；后续仍不默认进入 production app。真实产品验证前需明确用户、指标、试用人数、成功标准、失败标准和合规边界。",), ()),
    ]


def manual_sections() -> list[DocSection]:
    return [
        DocSection("1. 这个系统是做什么的", ("Makeup Engine 是模板生产和试用准备系统，可在浏览器本地运行，用于准备 PWA / Mobile Web MVP 和匿名内部试用；当前不是正式 App。",), (), "project-position"),
        DocSection("2. 如何启动项目", ("在 MacBook 上运行：cd /Users/star/Makeup-Engine；npm install；npm run dev。浏览器打开终端显示的本地地址，例如 http://localhost:5173/。验证命令包括 npm run project:status、npm run project:context、npm run typecheck、npm run build、npm run test、npm run user-app:browser-qa -- --json。",), (), "startup-flow"),
        DocSection("3. 推荐日常工作目录", ("固定使用 /Users/star/Makeup-Engine，不要主要依赖 Codex worktree 临时目录。",), ()),
        DocSection("4. 分支工作流", ("每个 Phase 新建分支，完成后 commit + push GitHub；不直接在 main 上堆功能；不 git add -A；不提交 dist、node_modules、tmp、public/mediapipe、.DS_Store。当前 GitHub 仓库：https://github.com/Shidawow/Makeup-Engine。",), (), "branch-workflow"),
        DocSection("5. 如何查看当前项目状态", ("查看 START_HERE.md、docs/status/CURRENT_PROJECT_STATUS.md、docs/status/NEXT_ACTION.md、project-state/project-state.snapshot.json、project-state/latest-handoff.json。命令包括 npm run project:status、npm run project:context、node scripts/project-status.mjs --json、node scripts/context-pack.mjs --json。",), (), "status-recovery-flow"),
        DocSection("6. Template Studio 使用说明", ("Template Studio 入口包括 Source Image Intake、Vision Analysis、Overlay Debug、Editable Mask Workbench、Evidence Panel、Dataset Panel、Review Queue、Replay Viewer、Template Library、Publish Package、Package Preview、User App Prototype Consumer。",), (), "template-studio-wireframe"),
        DocSection("7. User App Shell 普通用户路径", ("普通用户路径包括发现妆容、查看妆容详情、准备工具、开始跟练、逐步查看步骤、查看本地进度、查看本地偏好、查看隐私说明。",), (), "ordinary-user-flow"),
        DocSection("8. User App Shell 管理员路径", ("管理员入口覆盖 PWA 检查、MVP 打磨、App 就绪度、移动端 QA、交互检查、MVP 试用包、反馈表预览、试用就绪度、模板内容 QA、试用模板选择、试用内容就绪度、MVP 发布就绪度、试用 Go/No-Go、内部试用运营、观察记录模板、试用结果复盘、试用结果复盘框架、问题分类汇总、下一步决策框架、试用迭代计划、迭代 backlog、优先级建议、试用学习总结、产品决策门、下一阶段建议、内部试用证据包、试用证据摘要、证据充分性判断、证据收集协议、证据收集 checklist、证据收集质量门、匿名内部试用 dry run、dry run checklist、dry run 复盘、匿名内部试用启动包、启动就绪度、试用后 handoff、匿名试用证据复盘、证据缺口复盘、下一步决策输入、匿名试用后续迭代、证据缺口行动计划、后续试用就绪度。",), (), "admin-overview", "user-vs-admin-path"),
        DocSection("9. 如何准备匿名内部试用", ("先确认 readiness gate，使用 9F evidence collection protocol、9G dry run pack、9H launch pack，确认 participant notice、admin script、allowed evidence、forbidden data、stop conditions。",), (), "trial-prep-flow"),
        DocSection("10. 如何复盘匿名内部试用", ("使用 9I evidence review、evidence gap review、decision input 和 9J follow-up iteration，判断是否继续匿名内部试用、重复 dry run、修 protocol、修 launch pack、暂停或准备 MVP validation preconditions。",), (), "trial-review-tree", "go-no-go-decision"),
        DocSection("11. 可以做什么 / 不可以做什么", ("可以做本地运行、查看模板、跟练流程、查看推荐、查看 PWA readiness、做内容 QA、准备匿名内部试用、做匿名证据复盘、做后续迭代计划。不可以做正式发布、公开招募、收集真实照片、收集真实姓名/手机/邮箱、收集健康信息/敏感身份/生物识别、上线 App Store、接后端、调 OpenAI API、训练模型、收集敏感资料。",), ()),
        DocSection("12. 常见问题", ("FAQ 覆盖：为什么不是 iOS App、为什么先做 PWA、能不能直接给用户用、能不能收集照片、能不能接 OpenAI、public/mediapipe 为什么没有、build chunk warning 是否严重、full test 怎么跑、Codex worktree 和 /Users/star/Makeup-Engine 的区别、如何确认代码已 push、为什么内部试用不等于正式发布、什么情况下必须暂停试用、为什么现在不继续堆 Phase 9K、为什么需要先补图文详细设计书和操作说明书。",), ()),
        DocSection("13. 后续工作", ("DOC-ILLUSTRATED 本次完成；后续可进入 Phase 10A MVP Validation Plan 或按匿名证据策略进入 Phase 9K。是否继续 production app discovery 应由匿名内部试用证据和明确产品验证指标决定。",), ()),
    ]


def create_docx(path: Path, title: str, subtitle: str, sections: Iterable[DocSection], diagram_map: dict[str, Diagram]) -> None:
    doc = Document()
    style_doc(doc, title)
    add_cover(doc, title, subtitle)
    add_callout(doc, "本次文档未采集真实界面截图，使用架构图、流程图和页面结构示意图替代。后续可在运行环境稳定后补充截图。")
    doc.add_heading("目录", level=1)
    for section in sections:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(section.title)
        set_run_font(run)
    for section in sections:
        add_section(doc, section, diagram_map)
    doc.save(path)


def write_delivery_note() -> None:
    DELIVERY_NOTE.write_text(
        "\n".join(
            [
                "# Makeup Engine 文档交付说明",
                "",
                "- 图文版详细设计书：/Users/star/Desktop/Makeup-Engine-图文版详细设计书.docx",
                "- 图文版操作说明书：/Users/star/Desktop/Makeup-Engine-图文版操作说明书.docx",
                "- 仓库 Markdown 源文档：docs/reports/makeup-engine-illustrated-detailed-design.md",
                "- 仓库 Markdown 源文档：docs/manuals/makeup-engine-illustrated-operation-manual.md",
                "- 图表源文件：docs/assets/diagrams/ 与 docs/assets/manual/",
                "",
                "说明：本次文档未采集真实界面截图，使用架构图、流程图和页面结构示意图替代。",
            ]
        )
        + "\n",
        encoding="utf-8",
    )


def main() -> None:
    ensure_dirs()
    diagram_map = diagrams()
    generate_diagrams(diagram_map)
    REPORT_MD.write_text(design_markdown(diagram_map), encoding="utf-8")
    MANUAL_MD.write_text(manual_markdown(diagram_map), encoding="utf-8")
    create_docx(
        REPORT_DOCX,
        "Makeup Engine 图文版详细设计书",
        "Vision Pipeline、Template Studio、User App Shell、匿名内部试用体系与 Phase 9J 后续迭代",
        design_sections(),
        diagram_map,
    )
    create_docx(
        MANUAL_DOCX,
        "Makeup Engine 图文版操作说明书",
        "面向项目老板、产品负责人和内部试用管理员的本地运行与试用准备手册",
        manual_sections(),
        diagram_map,
    )
    write_delivery_note()
    print(REPORT_MD)
    print(MANUAL_MD)
    print(REPORT_DOCX)
    print(MANUAL_DOCX)
    print(DELIVERY_NOTE)


if __name__ == "__main__":
    main()
